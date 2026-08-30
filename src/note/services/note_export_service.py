from __future__ import annotations

import csv
from io import BytesIO
from typing import TYPE_CHECKING

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import strip_tags
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


if TYPE_CHECKING:
    from django.db.models import QuerySet

    from note.models import NoteProxy


class NoteExportService:
    """Service to export Note records to CSV and XLSX formats."""

    HEADERS = [
        "Title",
        "Section",
        "Client",
        "Job",
        "Task",
        "Created At",
        "Updated At",
        "Content",
    ]

    def export_csv(
        self, notes: QuerySet[NoteProxy] | list[NoteProxy]
    ) -> HttpResponse:
        """Export notes to CSV file format."""
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response["Content-Disposition"] = (
            f'attachment; filename="notes_export_{timestamp}.csv"'
        )

        writer = csv.writer(response)
        writer.writerow(self.HEADERS)

        notes_iterable = (
            notes.iterator(chunk_size=500)
            if hasattr(notes, "iterator")
            else notes
        )
        for note in notes_iterable:
            client_name = note.client.name if note.client else ""
            job_title = note.job.title if note.job else ""
            task_title = note.task.title if note.task else ""
            section_display = note.get_note_section_display() if hasattr(note, "get_note_section_display") else (note.note_section or "")
            clean_body = strip_tags(note.body or "").strip()

            writer.writerow(
                [
                    note.title,
                    section_display,
                    client_name,
                    job_title,
                    task_title,
                    note.created_at.strftime("%Y-%m-%d %H:%M") if note.created_at else "",
                    note.updated_at.strftime("%Y-%m-%d %H:%M") if note.updated_at else "",
                    clean_body,
                ]
            )

        return response

    def export_xlsx(
        self, notes: QuerySet[NoteProxy] | list[NoteProxy]
    ) -> HttpResponse:
        """Export notes to styled Excel (.xlsx) file format."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Notes"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_side = Side(border_style="thin", color="E5E7EB")
        border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        zebra_fill = PatternFill(
            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
        )
        white_fill = PatternFill(
            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
        )

        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        notes_iterable = (
            notes.iterator(chunk_size=500)
            if hasattr(notes, "iterator")
            else notes
        )
        for row_idx, note in enumerate(notes_iterable, start=2):
            client_name = note.client.name if note.client else ""
            job_title = note.job.title if note.job else ""
            task_title = note.task.title if note.task else ""
            section_display = note.get_note_section_display() if hasattr(note, "get_note_section_display") else (note.note_section or "")
            clean_body = strip_tags(note.body or "").strip()

            row_data = [
                note.title,
                section_display,
                client_name,
                job_title,
                task_title,
                note.created_at.strftime("%Y-%m-%d %H:%M") if note.created_at else "",
                note.updated_at.strftime("%Y-%m-%d %H:%M") if note.updated_at else "",
                clean_body,
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            row_fill = zebra_fill if row_idx % 2 == 0 else white_fill

            for col_num in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.fill = row_fill
                c.border = border
                c.font = Font(name="Calibri", size=10)
                if col_num in (2, 6, 7):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        column_widths = {
            "A": 28,  # Title
            "B": 16,  # Section
            "C": 24,  # Client
            "D": 24,  # Job
            "E": 24,  # Task
            "F": 18,  # Created At
            "G": 18,  # Updated At
            "H": 40,  # Content
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="notes_export_{timestamp}.xlsx"'
        )
        return response
