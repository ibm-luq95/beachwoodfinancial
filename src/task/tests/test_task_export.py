from __future__ import annotations

import io

import openpyxl
import pytest
from django.core.management import call_command
from django.test import RequestFactory
from django.urls import reverse
from model_bakery import baker

from beach_wood_user.models import BWUser
from client.models import ClientProxy
from core.choices import JobStatusEnum, TaskStatusEnum, TaskTypeEnum
from core.constants.users import CON_MANAGER
from job.models import JobProxy
from task.models import TaskProxy
from task.services import TaskExportService
from task.views import TaskExportCsvView, TaskExportExcelView, TaskExportView

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def setup_groups(db):
    """Ensure permission groups exist for tests."""
    call_command("create_groups")


def test_task_export_service_csv() -> None:
    client = ClientProxy.objects.create(name="Acme Corp")
    job = baker.make(
        JobProxy,
        title="Audit 2026",
        client=client,
        status=JobStatusEnum.IN_PROGRESS,
    )
    task = baker.make(
        TaskProxy,
        title="Prepare Ledger",
        job=job,
        status=TaskStatusEnum.IN_PROGRESS,
        task_type=TaskTypeEnum.URGENT,
        is_completed=False,
        hints="Check Q1 transactions",
    )

    service = TaskExportService()
    response = service.export_csv([task])

    assert response.status_code == 200
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert 'filename="tasks_export.csv"' in response["Content-Disposition"]

    content = response.content.decode("utf-8")
    assert "Task Title" in content
    assert "Job Title" in content
    assert "Prepare Ledger" in content
    assert "Audit 2026" in content
    assert "Acme Corp" in content
    assert "Check Q1 transactions" in content


def test_task_export_service_xlsx() -> None:
    client = ClientProxy.objects.create(name="Beta LLC")
    job = baker.make(
        JobProxy,
        title="Tax Return 2026",
        client=client,
        status=JobStatusEnum.IN_PROGRESS,
    )
    task = baker.make(
        TaskProxy,
        title="File Schedule C",
        job=job,
        status=TaskStatusEnum.COMPLETED,
        task_type=TaskTypeEnum.ONE_TIME,
        is_completed=True,
    )

    service = TaskExportService()
    response = service.export_xlsx([task])

    assert response.status_code == 200
    assert (
        response["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="tasks_export.xlsx"' in response["Content-Disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Tasks" in wb.sheetnames
    ws = wb["Tasks"]
    assert ws.cell(row=1, column=1).value == "Task Title"
    assert ws.cell(row=2, column=1).value == "File Schedule C"
    assert ws.cell(row=2, column=2).value == "Tax Return 2026"
    assert ws.cell(row=2, column=3).value == "Beta LLC"
    assert ws.cell(row=2, column=6).value == "Yes"


def test_task_export_csv_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(TaskProxy, title="Exportable Task CSV")

    factory = RequestFactory()
    request = factory.get(reverse("dashboard:task:export_csv"))
    request.user = admin_user

    view = TaskExportCsvView()
    view.setup(request)
    response = view.get(request)

    assert response.status_code == 200
    assert "Exportable Task CSV" in response.content.decode("utf-8")


def test_task_export_excel_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(TaskProxy, title="Exportable Task Excel")

    factory = RequestFactory()
    request = factory.get(reverse("dashboard:task:export_excel"))
    request.user = admin_user

    view = TaskExportExcelView()
    view.setup(request)
    response = view.get(request)

    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb["Tasks"]
    assert ws.cell(row=2, column=1).value == "Exportable Task Excel"


def test_task_export_unified_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(TaskProxy, title="Unified Task Export")

    factory = RequestFactory()

    # CSV format test
    req_csv = factory.get(f"{reverse('dashboard:task:export')}?format=csv")
    req_csv.user = admin_user
    view_csv = TaskExportView()
    view_csv.setup(req_csv)
    res_csv = view_csv.get(req_csv)
    assert res_csv.status_code == 200
    assert "Unified Task Export" in res_csv.content.decode("utf-8")

    # Excel format test
    req_xlsx = factory.get(f"{reverse('dashboard:task:export')}?format=xlsx")
    req_xlsx.user = admin_user
    view_xlsx = TaskExportView()
    view_xlsx.setup(req_xlsx)
    res_xlsx = view_xlsx.get(req_xlsx)
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    ws = wb["Tasks"]
    assert ws.cell(row=2, column=1).value == "Unified Task Export"


def test_task_export_view_filter_scoping(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(
        TaskProxy,
        title="Active Task Item",
        status=TaskStatusEnum.IN_PROGRESS,
    )
    baker.make(
        TaskProxy,
        title="Finished Task Item",
        status=TaskStatusEnum.COMPLETED,
    )

    factory = RequestFactory()
    request = factory.get(
        f"{reverse('dashboard:task:export_csv')}?status={TaskStatusEnum.IN_PROGRESS}"
    )
    request.user = admin_user

    view = TaskExportCsvView()
    view.setup(request)
    response = view.get(request)

    content = response.content.decode("utf-8")
    assert "Active Task Item" in content
    assert "Finished Task Item" not in content
