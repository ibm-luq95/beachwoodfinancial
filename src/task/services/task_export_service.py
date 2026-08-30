"""Service to handle CSV and XLSX export generation for Tasks."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from typing import TYPE_CHECKING, Any, ClassVar

import openpyxl
from django.http import HttpResponse
from django.utils.translation import gettext as _
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


if TYPE_CHECKING:
    from task.models import TaskProxy


class TaskExportService:
    """Service to generate exported CSV or XLSX files from a Task queryset."""

    HEADERS: ClassVar[list[str]] = [
        "Task Title",
        "Job Title",
        "Client Name",
        "Task Type",
        "Status",
        "Completed",
        "Job Manager",
        "Job Due Date",
        "Hints",
        "Created At",
        "Updated At",
    ]

    def _get_row_data(self, task: TaskProxy) -> list[Any]:
        """Convert a task instance into an export row."""
        job = getattr(task, "job", None)
        client = getattr(job, "client", None) if job else None
        manager = getattr(job, "managed_by", None) if job else None

        return [
            task.title or "",
            job.title if job else "",
            str(client.name) if client else "",
            task.get_task_type_display() if task.task_type else "",
            task.get_status_display() if task.status else "",
            "Yes" if task.is_completed else "No",
            str(manager) if manager else "",
            job.due_date.strftime("%Y-%m-%d") if (job and job.due_date) else "",
            task.hints or "",
            task.created_at.strftime("%Y-%m-%d %H:%M") if task.created_at else "",
            task.updated_at.strftime("%Y-%m-%d %H:%M") if task.updated_at else "",
        ]

    def export_csv(self, tasks: Iterable[TaskProxy]) -> HttpResponse:
        """Export tasks queryset into a downloadable CSV HttpResponse."""
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="tasks_export.csv"'

        writer = csv.writer(response)
        writer.writerow(self.HEADERS)
        for task in tasks:
            writer.writerow(self._get_row_data(task))
        return response

    def export_xlsx(self, tasks: Iterable[TaskProxy]) -> HttpResponse:
        """Export tasks queryset into a styled XLSX HttpResponse."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tasks"

        # Header Styles
        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )
        data_font = Font(name="Calibri", size=10)
        data_alignment = Alignment(vertical="center")
        zebra_fill = PatternFill(
            start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
        )

        # Write header
        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28
        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data rows
        row_idx = 2
        for task in tasks:
            row_data = self._get_row_data(task)
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = zebra_fill
            row_idx += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="tasks_export.xlsx"'
        return response
