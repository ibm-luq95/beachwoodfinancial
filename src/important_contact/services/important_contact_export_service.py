from __future__ import annotations

import csv
from io import BytesIO
from typing import TYPE_CHECKING

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


if TYPE_CHECKING:
    from django.db.models import QuerySet

    from important_contact.models import ImportantContactProxy


class ImportantContactExportService:
    """Service to export ImportantContact records to CSV and XLSX formats."""

    HEADERS = [
        "Full Name",
        "Role / Label",
        "Company Name",
        "Email",
        "Phone",
        "City",
        "State",
        "Postcode",
        "Website",
        "Linked Clients",
        "Created At",
        "Notes",
    ]

    def export_csv(
        self, contacts: QuerySet[ImportantContactProxy] | list[ImportantContactProxy]
    ) -> HttpResponse:
        """Export contacts to CSV file format."""
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response["Content-Disposition"] = (
            f'attachment; filename="important_contacts_export_{timestamp}.csv"'
        )

        writer = csv.writer(response)
        writer.writerow(self.HEADERS)

        contacts_iterable = (
            contacts.iterator(chunk_size=500)
            if hasattr(contacts, "iterator")
            else contacts
        )

        for contact in contacts_iterable:
            clients_list = ", ".join([c.name for c in contact.client.all() if c.name])
            role_display = (
                contact.get_contact_label_display()
                if hasattr(contact, "get_contact_label_display")
                else (contact.contact_label or "")
            )

            writer.writerow([
                contact.full_name,
                role_display,
                contact.company_name or "",
                contact.contact_email or "",
                contact.contact_phone or "",
                contact.contact_city or "",
                contact.contact_state or "",
                contact.contact_postcode or "",
                contact.contact_website or "",
                clients_list,
                (
                    contact.created_at.strftime("%Y-%m-%d %H:%M")
                    if contact.created_at
                    else ""
                ),
                contact.contact_notes or "",
            ])

        return response

    def export_xlsx(
        self, contacts: QuerySet[ImportantContactProxy] | list[ImportantContactProxy]
    ) -> HttpResponse:
        """Export contacts to styled Excel (.xlsx) file format."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Important Contacts"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        thin_side = Side(border_style="thin", color="E5E7EB")
        border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        zebra_fill = PatternFill(
            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
        )
        white_fill = PatternFill(
            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
        )

        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        contacts_iterable = (
            contacts.iterator(chunk_size=500)
            if hasattr(contacts, "iterator")
            else contacts
        )

        for row_idx, contact in enumerate(contacts_iterable, start=2):
            clients_list = ", ".join([c.name for c in contact.client.all() if c.name])
            role_display = (
                contact.get_contact_label_display()
                if hasattr(contact, "get_contact_label_display")
                else (contact.contact_label or "")
            )

            row_data = [
                contact.full_name,
                role_display,
                contact.company_name or "",
                contact.contact_email or "",
                contact.contact_phone or "",
                contact.contact_city or "",
                contact.contact_state or "",
                contact.contact_postcode or "",
                contact.contact_website or "",
                clients_list,
                (
                    contact.created_at.strftime("%Y-%m-%d %H:%M")
                    if contact.created_at
                    else ""
                ),
                contact.contact_notes or "",
            ]
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            row_fill = zebra_fill if row_idx % 2 == 0 else white_fill

            for col_num in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.fill = row_fill
                c.border = border
                c.font = Font(name="Calibri", size=10)
                if col_num in (2, 6, 7, 8, 11):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        column_widths = {
            "A": 24,  # Full Name
            "B": 16,  # Role
            "C": 24,  # Company Name
            "D": 28,  # Email
            "E": 20,  # Phone
            "F": 16,  # City
            "G": 12,  # State
            "H": 12,  # Postcode
            "I": 24,  # Website
            "J": 28,  # Linked Clients
            "K": 18,  # Created At
            "L": 35,  # Notes
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="important_contacts_export_{timestamp}.xlsx"'
        )
        return response
