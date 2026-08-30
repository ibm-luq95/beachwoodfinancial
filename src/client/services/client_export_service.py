"""Client export service module for exporting client records to CSV and XLSX formats."""

from __future__ import annotations

import csv
from io import BytesIO
from typing import TYPE_CHECKING, ClassVar

import openpyxl
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


if TYPE_CHECKING:
    from django.db.models import QuerySet

    from client.models import ClientProxy


class ClientExportService:
    """Service to export Client records to CSV and XLSX formats."""

    HEADERS: ClassVar[list[str]] = [
        "Client Name",
        "Industry",
        "Status",
        "Email",
        "Phone Number",
        "City",
        "State",
        "Managed By (Bookkeepers)",
        "CFOs",
        "Total Jobs",
        "Past Due Jobs",
        "Created At",
    ]

    def _extract_row_data(self, client_obj: ClientProxy) -> list[str | int]:
        """Extract serializable row data from a ClientProxy instance."""
        bookkeepers_str = ", ".join(str(b) for b in client_obj.bookkeepers.all())
        cfos_str = ", ".join(str(c) for c in client_obj.cfos.all())
        status_display = (
            client_obj.get_status_display()
            if hasattr(client_obj, "get_status_display")
            else (client_obj.status or "")
        )
        created_str = (
            client_obj.created_at.strftime("%Y-%m-%d") if client_obj.created_at else ""
        )

        return [
            client_obj.name or "",
            client_obj.industry or "",
            status_display,
            client_obj.email or "",
            client_obj.phone_number or "",
            client_obj.city or "",
            client_obj.state or "",
            bookkeepers_str or "Unassigned",
            cfos_str or "None",
            getattr(client_obj, "total_jobs_count", 0),
            getattr(client_obj, "past_due_jobs_count", 0),
            created_str,
        ]

    def export_csv(
        self, clients: QuerySet[ClientProxy] | list[ClientProxy]
    ) -> HttpResponse:
        """Export clients to CSV file format."""
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response["Content-Disposition"] = (
            f'attachment; filename="clients_export_{timestamp}.csv"'
        )

        writer = csv.writer(response)
        writer.writerow(self.HEADERS)

        clients_iterable = (
            clients.iterator(chunk_size=500)
            if hasattr(clients, "iterator")
            else clients
        )

        for client_obj in clients_iterable:
            writer.writerow(self._extract_row_data(client_obj))

        return response

    def export_xlsx(
        self, clients: QuerySet[ClientProxy] | list[ClientProxy]
    ) -> HttpResponse:
        """Export clients to styled Excel (.xlsx) file format."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clients"
        ws.views.sheetView[0].showGridLines = True

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

        thin_side = Side(border_style="thin", color="E5E7EB")
        border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        zebra_fill = PatternFill(
            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
        )
        white_fill = PatternFill(
            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
        )

        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28

        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        clients_iterable = (
            clients.iterator(chunk_size=500)
            if hasattr(clients, "iterator")
            else clients
        )

        for row_idx, client_obj in enumerate(clients_iterable, start=2):
            row_data = self._extract_row_data(client_obj)
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            row_fill = zebra_fill if row_idx % 2 == 0 else white_fill

            for col_num in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.fill = row_fill
                c.border = border
                c.font = Font(name="Calibri", size=10)
                if col_num in (3, 6, 7, 10, 11, 12):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        column_widths = {
            "A": 26,  # Client Name
            "B": 18,  # Industry
            "C": 14,  # Status
            "D": 28,  # Email
            "E": 18,  # Phone Number
            "F": 16,  # City
            "G": 12,  # State
            "H": 28,  # Managed By
            "I": 20,  # CFOs
            "J": 14,  # Total Jobs
            "K": 14,  # Past Due Jobs
            "L": 16,  # Created At
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="clients_export_{timestamp}.xlsx"'
        )
        return response
