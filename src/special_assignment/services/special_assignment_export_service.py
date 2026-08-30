from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from typing import TYPE_CHECKING, Any, ClassVar

import openpyxl
from django.http import HttpResponse
from django.utils.html import strip_tags
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from special_assignment.models import SpecialAssignmentProxy


class SpecialAssignmentExportService:
    """Service to generate exported CSV or XLSX files from a SpecialAssignment queryset."""

    HEADERS: ClassVar[list[str]] = [
        "Title",
        "Status",
        "Seen",
        "Client",
        "Job",
        "Assigned To",
        "Assigned By",
        "Start Date",
        "Due Date",
        "Has Attachment",
        "Created At",
        "Updated At",
        "Instructions",
    ]

    def _get_row_data(self, assignment: SpecialAssignmentProxy) -> list[Any]:
        client = getattr(assignment, "client", None)
        job = getattr(assignment, "job", None)
        assigned_to = getattr(assignment, "assigned_to", None)
        assigned_by = getattr(assignment, "assigned_by", None)

        body_cleaned = strip_tags(assignment.body or "").replace("\n", " ").strip()

        return [
            assignment.title or "",
            assignment.get_status_display() if assignment.status else "",
            "Yes" if assignment.is_seen else "No",
            client.name if client else "Standalone",
            job.title if job else "",
            (
                assigned_to.fullname
                if assigned_to and hasattr(assigned_to, "fullname")
                else str(assigned_to or "Unassigned")
            ),
            (
                assigned_by.fullname
                if assigned_by and hasattr(assigned_by, "fullname")
                else str(assigned_by or "System")
            ),
            (
                assignment.start_date.strftime("%Y-%m-%d")
                if assignment.start_date
                else ""
            ),
            (assignment.due_date.strftime("%Y-%m-%d") if assignment.due_date else ""),
            "Yes" if assignment.attachment else "No",
            (
                assignment.created_at.strftime("%Y-%m-%d %H:%M")
                if assignment.created_at
                else ""
            ),
            (
                assignment.updated_at.strftime("%Y-%m-%d %H:%M")
                if assignment.updated_at
                else ""
            ),
            body_cleaned[:250],
        ]

    def export_csv(self, assignments: Iterable[SpecialAssignmentProxy]) -> HttpResponse:
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = (
            'attachment; filename="special_assignments_export.csv"'
        )

        writer = csv.writer(response)
        writer.writerow(self.HEADERS)
        for assignment in assignments:
            writer.writerow(self._get_row_data(assignment))
        return response

    def export_xlsx(
        self, assignments: Iterable[SpecialAssignmentProxy]
    ) -> HttpResponse:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Special Assignments"

        header_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )
        data_font = Font(name="Calibri", size=10)
        data_alignment = Alignment(vertical="center")
        zebra_fill = PatternFill(
            start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
        )

        ws.append(self.HEADERS)
        ws.row_dimensions[1].height = 28
        for col_num in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        row_idx = 2
        for assignment in assignments:
            row_data = self._get_row_data(assignment)
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = zebra_fill
            row_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = (
            'attachment; filename="special_assignments_export.xlsx"'
        )
        return response
