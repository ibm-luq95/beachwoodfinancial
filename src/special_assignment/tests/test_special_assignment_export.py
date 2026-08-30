# -*- coding: utf-8 -*-#
from __future__ import annotations

import io
from typing import Any
import openpyxl
import pytest
from django.contrib.sites.models import Site
from django.core.management import call_command
from django.test import RequestFactory
from django.urls import reverse
from model_bakery import baker

from beach_wood_user.models import BWUser
from client.models import ClientProxy
from core.choices.special_assignment import SpecialAssignmentStatusEnum
from core.constants.site_settings import SITE_SETTINGS_DB_SLUG
from core.constants.users import CON_MANAGER
from job.models import JobProxy
from site_settings.models import SiteSettings
from special_assignment.models import SpecialAssignmentProxy
from special_assignment.services import SpecialAssignmentExportService
from special_assignment.views import (
    SpecialAssignmentExportCsvView,
    SpecialAssignmentExportExcelView,
    SpecialAssignmentExportView,
)

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def setup_test_environment(db: Any) -> None:
    """Ensure Django auth groups, site settings, and sites exist before running tests."""
    call_command("create_groups")
    SiteSettings.objects.get_or_create(
        slug=SITE_SETTINGS_DB_SLUG,
        defaults={"name": "LedgerFlare", "email": "info@ledgerflare.com"},
    )
    Site.objects.get_or_create(domain="testserver", defaults={"name": "testserver"})


def test_special_assignment_export_service_csv(admin_user: BWUser) -> None:
    client = ClientProxy.objects.create(name="Delta Corp")
    job = baker.make(JobProxy, title="Audit 2026", client=client)
    assignment = baker.make(
        SpecialAssignmentProxy,
        title="Annual Tax Review",
        client=client,
        job=job,
        status=SpecialAssignmentStatusEnum.IN_PROGRESS,
        body="Verify all 2025 deductions.",
        assigned_to=admin_user,
        is_deleted=False,
    )

    service = SpecialAssignmentExportService()
    response = service.export_csv([assignment])

    assert response.status_code == 200
    assert response["Content-Type"] == "text/csv; charset=utf-8"
    assert (
        'filename="special_assignments_export.csv"' in response["Content-Disposition"]
    )

    content = response.content.decode("utf-8")
    assert "Title" in content
    assert "Annual Tax Review" in content
    assert "Delta Corp" in content
    assert "Audit 2026" in content
    assert "Verify all 2025 deductions." in content


def test_special_assignment_export_service_xlsx(admin_user: BWUser) -> None:
    client = ClientProxy.objects.create(name="Epsilon LLC")
    assignment = baker.make(
        SpecialAssignmentProxy,
        title="Quarterly Review Assignment",
        client=client,
        status=SpecialAssignmentStatusEnum.IN_PROGRESS,
        body="Complete quarterly metrics.",
        assigned_to=admin_user,
        is_deleted=False,
    )

    service = SpecialAssignmentExportService()
    response = service.export_xlsx([assignment])

    assert response.status_code == 200
    assert (
        response["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="special_assignments_export.xlsx"' in response["Content-Disposition"]
    )

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Special Assignments" in wb.sheetnames
    ws = wb["Special Assignments"]
    assert ws.cell(row=1, column=1).value == "Title"
    assert ws.cell(row=2, column=1).value == "Quarterly Review Assignment"
    assert ws.cell(row=2, column=4).value == "Epsilon LLC"


def test_special_assignment_export_csv_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(
        SpecialAssignmentProxy,
        title="Exportable SA CSV",
        assigned_to=admin_user,
        is_deleted=False,
    )

    factory = RequestFactory()
    request = factory.get(reverse("dashboard:special_assignment:export_csv"))
    request.user = admin_user

    view = SpecialAssignmentExportCsvView()
    view.setup(request)
    response = view.get(request)

    assert response.status_code == 200
    assert "Exportable SA CSV" in response.content.decode("utf-8")


def test_special_assignment_export_excel_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(
        SpecialAssignmentProxy,
        title="Exportable SA Excel",
        assigned_to=admin_user,
        is_deleted=False,
    )

    factory = RequestFactory()
    request = factory.get(reverse("dashboard:special_assignment:export_excel"))
    request.user = admin_user

    view = SpecialAssignmentExportExcelView()
    view.setup(request)
    response = view.get(request)

    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb["Special Assignments"]
    assert ws.cell(row=2, column=1).value == "Exportable SA Excel"


def test_special_assignment_export_unified_view(admin_user: BWUser) -> None:
    admin_user.user_type = CON_MANAGER
    admin_user.save()

    baker.make(
        SpecialAssignmentProxy,
        title="Unified SA Export Item",
        assigned_to=admin_user,
        is_deleted=False,
    )

    factory = RequestFactory()

    # Test CSV
    req_csv = factory.get(
        f"{reverse('dashboard:special_assignment:export')}?format=csv"
    )
    req_csv.user = admin_user
    view_csv = SpecialAssignmentExportView()
    view_csv.setup(req_csv)
    res_csv = view_csv.get(req_csv)
    assert res_csv.status_code == 200
    assert "Unified SA Export Item" in res_csv.content.decode("utf-8")

    # Test XLSX
    req_xlsx = factory.get(
        f"{reverse('dashboard:special_assignment:export')}?format=xlsx"
    )
    req_xlsx.user = admin_user
    view_xlsx = SpecialAssignmentExportView()
    view_xlsx.setup(req_xlsx)
    res_xlsx = view_xlsx.get(req_xlsx)
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    ws = wb["Special Assignments"]
    assert ws.cell(row=2, column=1).value == "Unified SA Export Item"
