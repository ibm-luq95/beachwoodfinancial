"""Jobs report export service unit tests."""
from __future__ import annotations

import io
import pytest
from model_bakery import baker
import openpyxl

from client.models import ClientProxy
from core.choices.job import JobStatusEnum
from reports.services.jobs_report_export_service import JobsReportExportService
from reports.services.jobs_report_service import JobsReportService


@pytest.mark.django_db
def test_export_excel_generates_valid_workbook() -> None:
    """Verify Excel export creates valid workbook with expected sheets and headers."""
    client = baker.make("client.Client", name="Acme Logistics", is_deleted=False)
    client_proxy = ClientProxy.objects.get(pk=client.pk)

    baker.make(
        "job.Job",
        client=client_proxy,
        period_year="2026",
        period_month="5",
        status=JobStatusEnum.COMPLETED,
        is_deleted=False,
    )

    report_service = JobsReportService()
    report_result = report_service.get_report_data(
        filter_params={"period_year": "2026", "clients": [client.pk]},
        page=1,
        per_page=10,
    )

    export_service = JobsReportExportService()
    buffer = export_service.export_excel(report_result=report_result, year="2026")

    assert isinstance(buffer, io.BytesIO)
    wb = openpyxl.load_workbook(buffer)
    assert "Jobs Report 2026" in wb.sheetnames

    ws = wb["Jobs Report 2026"]
    # Check title row
    assert "LedgerFlare — Client Jobs Annual Report" in str(ws["A1"].value)
    # Check table headers on row 7
    assert ws["A7"].value == "Client Name"
    assert ws["C7"].value == "Jan"
    assert ws["G7"].value == "May"
    assert ws["O7"].value == "Total Jobs"

    # Check client row data on row 8
    assert ws["A8"].value == "Acme Logistics"
    assert ws["G8"].value == 1  # May jobs
    assert ws["O8"].value == 1  # Total jobs


@pytest.mark.django_db
def test_export_csv_generates_valid_content() -> None:
    """Verify CSV export produces valid comma-separated text."""
    client = baker.make("client.Client", name="Beta Corp", is_deleted=False)
    client_proxy = ClientProxy.objects.get(pk=client.pk)

    baker.make(
        "job.Job",
        client=client_proxy,
        period_year="2026",
        period_month="3",
        status=JobStatusEnum.IN_PROGRESS,
        is_deleted=False,
    )

    report_service = JobsReportService()
    report_result = report_service.get_report_data(
        filter_params={"period_year": "2026", "clients": [client.pk]},
        page=1,
        per_page=10,
    )

    export_service = JobsReportExportService()
    csv_text = export_service.export_csv(report_result=report_result, year="2026")

    assert "Client Name,Categories,Jan,Feb,Mar" in csv_text
    assert "Beta Corp" in csv_text
    assert ",1,0,0%" in csv_text or "1,0" in csv_text
