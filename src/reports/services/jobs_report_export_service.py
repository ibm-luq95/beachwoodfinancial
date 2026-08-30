"""Jobs report export service for Excel and CSV outputs."""
from __future__ import annotations

import csv
import io
from typing import Any

import openpyxl
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reports.services.jobs_report_service import JobsReportResult


class JobsReportExportService:
    """Service to export client jobs reports into Excel and CSV formats."""

    def export_excel(
        self, report_result: JobsReportResult, year: str
    ) -> io.BytesIO:
        """Generate a professionally styled Excel workbook from report data."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Jobs Report {year}"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        navy_fill = PatternFill(
            start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
        )
        header_fill = PatternFill(
            start_color="2563EB", end_color="2563EB", fill_type="solid"
        )
        accent_fill = PatternFill(
            start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"
        )
        kpi_fill = PatternFill(
            start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
        )

        title_font = Font(name="Arial", size=15, bold=True, color="FFFFFF")
        sub_font = Font(name="Arial", size=9, italic=True, color="E2E8F0")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        bold_font = Font(name="Arial", size=9, bold=True)
        regular_font = Font(name="Arial", size=9)
        kpi_value_font = Font(name="Arial", size=13, bold=True, color="1E3A8A")

        thin_side = Side(style="thin", color="CBD5E1")
        cell_border = Border(
            left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
        )

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # 1. Title Banner (Rows 1-2)
        ws.merge_cells("A1:Q1")
        title_cell = ws["A1"]
        title_cell.value = f"LedgerFlare — Client Jobs Annual Report (FY {year})"
        title_cell.font = title_font
        title_cell.fill = navy_fill
        title_cell.alignment = left_align

        ws.merge_cells("A2:Q2")
        meta_cell = ws["A2"]
        meta_cell.value = (
            f"Generated on {timezone.now().strftime('%B %d, %Y at %H:%M UTC')} "
            f"| Total Clients: {report_result.total_clients_count} "
            f"| Total Jobs: {report_result.summary_kpis.total_jobs_count}"
        )
        meta_cell.font = sub_font
        meta_cell.fill = navy_fill
        meta_cell.alignment = left_align

        ws.row_dimensions[1].height = 26
        ws.row_dimensions[2].height = 18

        # 2. Executive KPI Cards Section (Rows 4-5)
        ws["A4"] = "TOTAL JOBS"
        ws["A5"] = report_result.summary_kpis.total_jobs_count
        ws["C4"] = "COMPLETED"
        ws["C5"] = (
            f"{report_result.summary_kpis.total_completed_count} "
            f"({report_result.summary_kpis.overall_completion_rate}%)"
        )
        ws["E4"] = "IN PROGRESS"
        ws["E5"] = report_result.summary_kpis.total_in_progress_count
        ws["G4"] = "PAST DUE"
        ws["G5"] = report_result.summary_kpis.total_past_due_count

        for col_idx in ["A", "C", "E", "G"]:
            ws[f"{col_idx}4"].font = bold_font
            ws[f"{col_idx}4"].fill = kpi_fill
            ws[f"{col_idx}5"].font = kpi_value_font
            ws[f"{col_idx}5"].fill = kpi_fill

        ws.row_dimensions[4].height = 16
        ws.row_dimensions[5].height = 22

        # 3. Table Column Headers (Row 7)
        headers = [
            "Client Name",
            "Categories",
            "Jan",
            "Feb",
            "Mar",
            "Apr",
            "May",
            "Jun",
            "Jul",
            "Aug",
            "Sep",
            "Oct",
            "Nov",
            "Dec",
            "Total Jobs",
            "Completed",
            "Completion %",
        ]

        start_row = 7
        ws.row_dimensions[start_row].height = 24

        for col_idx, header_text in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx)
            cell.value = header_text
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = (
                center_align if col_idx > 2 else left_align
            )
            cell.border = cell_border

        # 4. Data Rows
        current_row = start_row + 1
        for row_dto in report_result.rows:
            ws.row_dimensions[current_row].height = 20

            # Client Name
            c1 = ws.cell(row=current_row, column=1, value=row_dto.client.name)
            c1.font = bold_font
            c1.alignment = left_align
            c1.border = cell_border

            # Categories
            categories_str = ", ".join(
                c.name for c in row_dto.client.categories.all()
            )
            c2 = ws.cell(row=current_row, column=2, value=categories_str)
            c2.font = regular_font
            c2.alignment = left_align
            c2.border = cell_border

            # 12 Months
            for m_idx, month_cell in enumerate(row_dto.months):
                cell = ws.cell(row=current_row, column=3 + m_idx)
                if month_cell.has_jobs:
                    cell.value = month_cell.total_count
                    cell.font = bold_font
                else:
                    cell.value = "-"
                    cell.font = regular_font
                cell.alignment = center_align
                cell.border = cell_border

            # Total Jobs
            c15 = ws.cell(
                row=current_row, column=15, value=row_dto.total_jobs
            )
            c15.font = bold_font
            c15.alignment = center_align
            c15.border = cell_border

            # Completed Jobs
            c16 = ws.cell(
                row=current_row, column=16, value=row_dto.completed_jobs
            )
            c16.font = regular_font
            c16.alignment = center_align
            c16.border = cell_border

            # Completion %
            c17 = ws.cell(
                row=current_row,
                column=17,
                value=f"{row_dto.completion_rate}%",
            )
            c17.font = bold_font
            c17.alignment = right_align
            c17.border = cell_border

            current_row += 1

        # 5. Total Row
        if report_result.rows:
            ws.row_dimensions[current_row].height = 22
            total_label = ws.cell(row=current_row, column=1, value="TOTAL")
            total_label.font = bold_font
            total_label.fill = accent_fill
            total_label.border = cell_border

            ws.cell(row=current_row, column=2, value="").fill = accent_fill
            ws.cell(row=current_row, column=2).border = cell_border

            for col_i in range(3, 15):
                col_letter = get_column_letter(col_i)
                sum_cell = ws.cell(row=current_row, column=col_i)
                sum_cell.value = (
                    f"=SUM({col_letter}{start_row + 1}:{col_letter}{current_row - 1})"
                )
                sum_cell.font = bold_font
                sum_cell.fill = accent_fill
                sum_cell.alignment = center_align
                sum_cell.border = cell_border

            # Total Jobs sum
            c15 = ws.cell(row=current_row, column=15)
            c15.value = f"=SUM(O{start_row + 1}:O{current_row - 1})"
            c15.font = bold_font
            c15.fill = accent_fill
            c15.alignment = center_align
            c15.border = cell_border

            # Total Completed sum
            c16 = ws.cell(row=current_row, column=16)
            c16.value = f"=SUM(P{start_row + 1}:P{current_row - 1})"
            c16.font = bold_font
            c16.fill = accent_fill
            c16.alignment = center_align
            c16.border = cell_border

            # Overall rate
            c17 = ws.cell(
                row=current_row,
                column=17,
                value=f"{report_result.summary_kpis.overall_completion_rate}%",
            )
            c17.font = bold_font
            c17.fill = accent_fill
            c17.alignment = right_align
            c17.border = cell_border

        # 6. Auto-fit Column Widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 2 and cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 10)

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 20

        # Freeze Panes
        ws.freeze_panes = f"A{start_row + 1}"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def export_csv(
        self, report_result: JobsReportResult, year: str
    ) -> str:
        """Generate CSV string from report data."""
        output = io.StringIO()
        writer = csv.writer(output)

        # Header
        writer.writerow(
            [
                "Client Name",
                "Categories",
                "Jan",
                "Feb",
                "Mar",
                "Apr",
                "May",
                "Jun",
                "Jul",
                "Aug",
                "Sep",
                "Oct",
                "Nov",
                "Dec",
                "Total Jobs",
                "Completed Jobs",
                "Completion Rate",
            ]
        )

        for row in report_result.rows:
            categories_str = "; ".join(
                c.name for c in row.client.categories.all()
            )
            months_data = [
                m.total_count if m.has_jobs else 0 for m in row.months
            ]
            writer.writerow(
                [
                    row.client.name,
                    categories_str,
                    *months_data,
                    row.total_jobs,
                    row.completed_jobs,
                    f"{row.completion_rate}%",
                ]
            )

        return output.getvalue()
